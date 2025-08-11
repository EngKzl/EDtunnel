from pptx import Presentation
from pptx.util import Pt


def add_title_slide(prs: Presentation, title: str, subtitle: str = "") -> None:
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    if subtitle:
        slide.placeholders[1].text = subtitle


def add_bullet_slide(
    prs: Presentation,
    title: str,
    bullets: list[str] | list[tuple[str, int]]
) -> None:
    slide_layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title

    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()

    def add_paragraph(text: str, level: int = 0) -> None:
        p = text_frame.add_paragraph() if len(text_frame.paragraphs) > 0 else text_frame.paragraphs[0]
        p.text = text
        p.level = level
        for run in p.runs:
            run.font.size = Pt(20)

    for idx, item in enumerate(bullets):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        if idx == 0:
            # Reuse first paragraph slot
            text_frame.paragraphs[0].text = text
            text_frame.paragraphs[0].level = level
            for run in text_frame.paragraphs[0].runs:
                run.font.size = Pt(20)
        else:
            add_paragraph(text, level)


if __name__ == "__main__":
    prs = Presentation()

    # 1. Title
    add_title_slide(
        prs,
        title="PEMBELAJARAN MENDALAM",
        subtitle="Ringkasan Materi & Panduan Implementasi"
    )

    # 2. Gambaran Umum
    add_bullet_slide(
        prs,
        title="Gambaran Umum",
        bullets=[
            "PM diterapkan dari PAUD, pendidikan dasar, hingga menengah.",
            "Tahap awal diprioritaskan pada pendidikan dasar dan menengah.",
            "Tujuan: pendidikan yang bermakna, kontekstual, dan berpihak pada murid."
        ],
    )

    # 3-4. Karakteristik Kurikulum
    add_bullet_slide(
        prs,
        title="Karakteristik Kurikulum (1/2)",
        bullets=[
            ("Dinamis, Fleksibel, Responsif:", 0),
            ("Dapat diperbarui mengikuti teknologi, budaya, kebutuhan masyarakat.", 1),
            ("Fleksibel sesuai konteks lokal tanpa kehilangan relevansi global.", 1),
            ("Berpusat pada Peserta Didik:", 0),
            ("Memberi ruang personalisasi sesuai minat, motivasi, gaya belajar.", 1),
            ("Pembelajaran Terpadu:", 0),
            ("Lintas disiplin, multidisipliner dan antardisiplin, terhubung ke kehidupan nyata.", 1),
        ],
    )

    add_bullet_slide(
        prs,
        title="Karakteristik Kurikulum (2/2)",
        bullets=[
            ("Relevan & Peduli Masyarakat:", 0),
            ("Isu nyata: sosial, politik, kesehatan, energi, lingkungan; dorong kontribusi nyata.", 1),
            ("Keterampilan Tingkat Tinggi:", 0),
            ("Kreativitas, kolaborasi, berpikir kritis, pemecahan masalah via proyek dan penelitian.", 1),
            ("Pemanfaatan Teknologi Digital:", 0),
            ("Interaksi diperkuat platform daring/luring; jangkau daerah terpencil.", 1),
        ],
    )

    # 5. Perencanaan PM
    add_bullet_slide(
        prs,
        title="Perencanaan Pembelajaran Mendalam (PM)",
        bullets=[
            "1) Identifikasi",
            "2) Desain Pembelajaran",
            "3) Pengalaman Belajar",
            "4) Asesmen"
        ],
    )

    # 6. Identifikasi - Kesiapan Murid
    add_bullet_slide(
        prs,
        title="Identifikasi: Kesiapan Murid",
        bullets=[
            ("Analisis pengetahuan awal (pre-test, diskusi, pertanyaan pemantik).", 0),
            ("Observasi & refleksi: keterlibatan, rasa ingin tahu, koneksi dengan pengalaman.", 0),
            ("Inventarisasi gaya belajar & minat (angket/wawancara).", 0),
            ("Analisis HOTS melalui tugas/proyek kecil untuk pemetaan pendampingan.", 0),
            ("Perhatikan konteks sosial-emosional dan diferensiasi strategi.", 0),
        ],
    )

    # 7. Identifikasi - Karakteristik Mapel
    add_bullet_slide(
        prs,
        title="Identifikasi: Karakteristik Mata Pelajaran",
        bullets=[
            ("Fokus pada 6C: Character, Citizenship, Collaboration, Communication, Creativity, Critical Thinking.", 0),
            ("Selaraskan dengan 8 Dimensi Profil Lulusan.", 0),
            ("Tentukan kompetensi esensial; hindari fokus hafalan semata.", 0),
            ("Kaitkan dengan konteks nyata (mis. IPA–lingkungan, Ekonomi–bisnis digital).", 0),
            ("Gunakan strategi eksploratif & reflektif: PjBL, PBL, inquiry, QBL.", 0),
        ],
    )

    # 8. Identifikasi - Dimensi Profil Lulusan
    add_bullet_slide(
        prs,
        title="Identifikasi: Dimensi Profil Lulusan",
        bullets=[
            "1) Keimanan & Ketakwaan",
            "2) Kewargaan",
            "3) Penalaran Kritis",
            "4) Kreativitas",
            "5) Kolaborasi",
            "6) Kemandirian",
            "7) Kesehatan",
            "8) Komunikasi",
        ],
    )

    # 9. Desain - Topik Kontekstual
    add_bullet_slide(
        prs,
        title="Desain: Topik Kontekstual & Relevan",
        bullets=[
            ("Pahami kebutuhan, minat, latar belakang, dan penguasaan awal murid.", 0),
            ("Hubungkan ke kehidupan nyata: studi kasus, isu aktual, isu sosial.", 0),
            ("Contoh topik: perubahan iklim, kesehatan masyarakat, literasi keuangan.", 0),
            ("Manfaatkan sumber digital: video interaktif, simulasi, eksperimen virtual.", 0),
        ],
    )

    # 10. Desain - Lintas Disiplin
    add_bullet_slide(
        prs,
        title="Desain: Lintas Disiplin",
        bullets=[
            ("Mulai dari tema utama; hubungkan konsep dari berbagai bidang.", 0),
            ("Gunakan STEAM/PBL untuk pemahaman holistik masalah nyata.", 0),
            ("Proyek lintas mapel untuk berpikir kritis & sistematis.", 0),
        ],
    )

    # 11. Desain - Kerangka PM
    add_bullet_slide(
        prs,
        title="Desain: Kerangka Pembelajaran Mendalam",
        bullets=[
            ("Praktik Pedagogis: proyek, inquiry, problem-based learning.", 0),
            ("Kemitraan Pembelajaran: orang tua, komunitas, industri, akademisi.", 0),
            ("Lingkungan Belajar: kelas kreatif, lab; forum LMS & ruang digital.", 0),
            ("Teknologi Digital: e-learning, simulasi, AR/AI untuk personalisasi.", 0),
        ],
    )

    # 12. Pengalaman Belajar - Prinsip
    add_bullet_slide(
        prs,
        title="Pengalaman Belajar: Prinsip",
        bullets=[
            ("Berkesadaran: perhatikan kebutuhan unik, keberagaman, kondisi nyata.", 0),
            ("Bermakna: dekat dengan pengalaman murid; bangun pemahaman mendalam.", 0),
            ("Menggembirakan: permainan edukatif, eksperimen interaktif, proyek kolaboratif.", 0),
        ],
    )

    # 13. Pengalaman Belajar - Tahapan
    add_bullet_slide(
        prs,
        title="Pengalaman Belajar: Tahapan",
        bullets=[
            ("Memahami: teori, diskusi aktif, demonstrasi.", 0),
            ("Mengaplikasikan: tugas bermakna, eksperimen, proyek kontekstual.", 0),
            ("Merefleksi: tinjau capaian, tantangan, strategi perbaikan.", 0),
        ],
    )

    # 14. Implementasi PM
    add_bullet_slide(
        prs,
        title="Implementasi PM",
        bullets=[
            ("Perencanaan: karakteristik murid, materi kontekstual, sumber daya, kolaborasi.", 0),
            ("Pelaksanaan: berkesadaran, bermakna, menggembirakan; memahami–mengaplikasi–merefleksi.", 0),
            ("Asesmen: kedalaman konsep, berpikir kritis, kesiapan aplikasi nyata.", 0),
        ],
    )

    # 15. Asesmen - Awal
    add_bullet_slide(
        prs,
        title="Asesmen: Awal Pembelajaran",
        bullets=[
            ("Tujuan: memetakan kesiapan, pengetahuan awal, kebutuhan individual.", 0),
            ("Manfaat: sesuaikan pendekatan, rancang pembelajaran inklusif, identifikasi kesenjangan.", 0),
            ("Metode: pre-test, diskusi awal, kuesioner, survei minat, wawancara singkat.", 0),
        ],
    )

    # 16. Asesmen - Formatif
    add_bullet_slide(
        prs,
        title="Asesmen: Proses/Formatif",
        bullets=[
            ("Tujuan: pantau perkembangan, beri umpan balik real-time, sesuaikan strategi.", 0),
            ("Teknik: observasi, refleksi, diskusi kelompok, kuis singkat, jurnal, pertanyaan terbuka.", 0),
        ],
    )

    # 17. Asesmen - Sumatif
    add_bullet_slide(
        prs,
        title="Asesmen: Akhir/Sumatif",
        bullets=[
            ("Evaluasi capaian kompetensi dan kualitas pemahaman & penerapan.", 0),
            ("Gunakan penilaian autentik untuk keterampilan abad 21.", 0),
            ("Metode: ujian, portofolio, proyek, presentasi, studi kasus, produk nyata.", 0),
        ],
    )

    # 18. Penutup
    add_bullet_slide(
        prs,
        title="Penutup",
        bullets=[
            ("Pembelajaran Mendalam memerdekakan murid sebagai pencipta makna.", 0),
            ("Dengan perencanaan berkesadaran, bermakna, menggembirakan, sekolah membangun masa depan cerah.", 0),
        ],
    )

    output_path = "/workspace/Pembelajaran_Mendalam.pptx"
    prs.save(output_path)
    print(f"Presentation generated: {output_path}")